from io import BytesIO
from fastapi import HTTPException, status, UploadFile
from pypdf import PdfReader
from docx import Document as DocxDocument

from app.models.document import Document
from app.repositories import document_repository
from app.schemas.document import DocumentOut


class DocumentService:
    MAX_SIZE = 8 * 1024 * 1024  # 8 MB

    async def extract_text(self, file: UploadFile) -> str:
        content = await file.read()
        if len(content) > self.MAX_SIZE:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File too large (max 8 MB)",
            )

        filename = (file.filename or "").lower()
        content_type = file.content_type or ""

        if content_type == "application/pdf" or filename.endswith(".pdf"):
            reader = PdfReader(BytesIO(content))
            parts = []
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    parts.append(text)
            return "\n".join(parts).strip()

        if (
            content_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            or filename.endswith(".docx")
        ):
            doc = DocxDocument(BytesIO(content))
            return "\n".join(p.text for p in doc.paragraphs if p.text).strip()

        if (
            content_type in ("text/plain", "text/markdown")
            or filename.endswith(".txt")
            or filename.endswith(".md")
        ):
            return content.decode("utf-8", errors="ignore").strip()

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Upload PDF, DOCX, or plain text.",
        )

    async def upload(self, user_id: str, file: UploadFile) -> DocumentOut:
        text = await self.extract_text(file)
        if not text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Could not extract text from file",
            )
        doc = Document(
            user_id=user_id,
            filename=file.filename or "upload",
            content=text,
        )
        document_repository.create(doc)
        preview = text[:300] + ("…" if len(text) > 300 else "")
        return DocumentOut(
            id=doc.id,
            filename=doc.filename,
            uploaded_at=doc.uploaded_at,
            preview=preview,
        )


document_service = DocumentService()
